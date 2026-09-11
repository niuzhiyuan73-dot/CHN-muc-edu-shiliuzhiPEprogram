"""
把 Node 单元测试的 JUnit 报告排版成 Excel 表格。

用法：
    1. node --test --test-reporter=junit --test-reporter-destination=tools/test-report.xml "tests/**/*.test.js"
    2. python tools/tests-to-xlsx.py tools/test-report.xml <输出.xlsx>

依赖 openpyxl。
"""

import os
import sys
import time
import xml.etree.ElementTree as ET

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

TITLE_FONT = Font(bold=True, size=14)
HEAD_FONT = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="4472C4")
NOTE_FONT = Font(size=10, color="595959")
SECTION_FONT = Font(bold=True, size=11)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")
WRAP = Alignment(vertical="top", wrap_text=True)


def write_table(ws, start_row, headers, rows, widths=None):
    """从 start_row 起写一张带表头的表，返回下一行的行号。"""
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col, value=text)
        cell.font = HEAD_FONT
        cell.fill = HEAD_FILL
        cell.alignment = CENTER
        cell.border = BORDER

    for offset, row in enumerate(rows, start=1):
        for col, value in enumerate(row, start=1):
            cell = ws.cell(row=start_row + offset, column=col, value=value)
            cell.border = BORDER
            if isinstance(value, (int, float)):
                cell.alignment = CENTER
            else:
                cell.alignment = Alignment(vertical="center")

    if widths:
        for col, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(col)].width = width

    return start_row + len(rows) + 2


def parse_report(xml_path):
    """
    把 JUnit XML 拆成逐条用例。返回 (用例列表, 运行信息)。

    每条用例字段：file（测试文件名）、name（用例名）、status（通过/失败/跳过）、ms（耗时毫秒）。
    """
    root = ET.parse(xml_path).getroot()
    cases = []
    for case in root.findall('.//testcase'):
        if case.find('failure') is not None or case.find('error') is not None:
            status = '失败'
        elif case.find('skipped') is not None:
            status = '跳过'
        else:
            status = '通过'

        raw_file = case.get('file') or ''
        filename = os.path.basename(raw_file.replace('\\', '/'))
        try:
            ms = round(float(case.get('time') or 0) * 1000, 3)
        except ValueError:
            ms = 0.0

        cases.append({
            'file': filename,
            'name': case.get('name') or '',
            'status': status,
            'ms': ms,
        })

    info = {
        'total': len(cases),
        'passed': sum(1 for c in cases if c['status'] == '通过'),
        'failed': sum(1 for c in cases if c['status'] == '失败'),
        'skipped': sum(1 for c in cases if c['status'] == '跳过'),
        'file_count': len({c['file'] for c in cases}),
        'total_ms': round(sum(c['ms'] for c in cases), 1),
    }
    return cases, info


def sheet_overview(wb, cases, info, xml_path, generated_at):
    ws = wb.create_sheet("说明")
    ws.column_dimensions["A"].width = 118

    lines = [
        ("中央民族大学匹克球俱乐部 · 单元测试记录", TITLE_FONT),
        ("", None),
        ("一、这份表是什么", SECTION_FONT),
        (
            "记录一次完整 npm test 运行的全部用例结果，逐条列出所属测试文件、用例名称、结果与耗时。"
            "这些是纯 Node 单元测试，覆盖算法层与数据层（ELO、回放、榜单、视图组装等），"
            "不需要微信开发者工具，秒级即可跑完。",
            None,
        ),
        ("", None),
        ("二、本次运行", SECTION_FONT),
        (
            f"共 {info['total']} 条用例，分布于 {info['file_count']} 个测试文件；"
            f"通过 {info['passed']} 条、失败 {info['failed']} 条、跳过 {info['skipped']} 条；"
            f"全部耗时合计约 {info['total_ms'] / 1000:.2f} 秒。",
            None,
        ),
        ("", None),
        ("三、怎么复现", SECTION_FONT),
        (
            '1. node --test --test-reporter=junit --test-reporter-destination=tools/test-report.xml "tests/**/*.test.js"',
            None,
        ),
        ("2. python tools/tests-to-xlsx.py tools/test-report.xml 单元测试记录.xlsx", None),
        (
            "「测试用例明细」按运行顺序逐条记录，「测试文件汇总」按文件统计用例数与耗时。",
            None,
        ),
        ("", None),
        (f"报告文件：{xml_path}", NOTE_FONT),
        (f"生成时间：{generated_at}", NOTE_FONT),
    ]

    for row, (text, font) in enumerate(lines, start=1):
        cell = ws.cell(row=row, column=1, value=text)
        cell.alignment = WRAP
        if font:
            cell.font = font

    return ws


def sheet_cases(wb, cases):
    ws = wb.create_sheet("测试用例明细")
    ws.cell(row=1, column=1, value="全部测试用例（按运行顺序）").font = TITLE_FONT

    headers = ["序号", "测试文件", "用例名称", "结果", "耗时（毫秒）"]
    rows = [
        [index, c['file'], c['name'], c['status'], c['ms']]
        for index, c in enumerate(cases, start=1)
    ]

    write_table(ws, 3, headers, rows, widths=[6, 22, 72, 8, 12])
    ws.freeze_panes = "A4"
    return ws


def sheet_files(wb, cases, info):
    ws = wb.create_sheet("测试文件汇总")
    ws.cell(row=1, column=1, value="按测试文件汇总").font = TITLE_FONT

    per_file = {}
    for case in cases:
        bucket = per_file.setdefault(
            case['file'],
            {'count': 0, 'passed': 0, 'failed': 0, 'skipped': 0, 'ms': 0.0},
        )
        bucket['count'] += 1
        bucket['ms'] += case['ms']
        if case['status'] == '通过':
            bucket['passed'] += 1
        elif case['status'] == '失败':
            bucket['failed'] += 1
        else:
            bucket['skipped'] += 1

    headers = ["测试文件", "用例数", "通过", "失败", "跳过", "耗时合计（毫秒）", "平均耗时（毫秒）"]
    rows = []
    for filename in sorted(per_file):
        bucket = per_file[filename]
        avg = round(bucket['ms'] / bucket['count'], 3) if bucket['count'] else 0
        rows.append([
            filename, bucket['count'], bucket['passed'], bucket['failed'],
            bucket['skipped'], round(bucket['ms'], 1), avg,
        ])
    rows.append([
        "合计", info['total'], info['passed'], info['failed'],
        info['skipped'], info['total_ms'],
        round(info['total_ms'] / info['total'], 3) if info['total'] else 0,
    ])

    write_table(ws, 3, headers, rows, widths=[22, 9, 8, 8, 8, 16, 16])
    return ws


def main():
    if len(sys.argv) < 3:
        raise SystemExit("用法：python tools/tests-to-xlsx.py <JUnit XML> <输出.xlsx>")

    xml_path, xlsx_path = sys.argv[1], sys.argv[2]
    cases, info = parse_report(xml_path)
    generated_at = time.strftime(
        "%Y-%m-%d %H:%M:%S", time.localtime(os.path.getmtime(xml_path))
    )

    wb = Workbook()
    wb.remove(wb.active)  # 去掉默认空表

    sheet_overview(wb, cases, info, xml_path, generated_at)
    sheet_cases(wb, cases)
    sheet_files(wb, cases, info)

    wb.save(xlsx_path)
    print(
        f"已写出 {xlsx_path}：{info['total']} 条用例，"
        f"通过 {info['passed']}，失败 {info['failed']}，跳过 {info['skipped']}"
    )


if __name__ == "__main__":
    main()
